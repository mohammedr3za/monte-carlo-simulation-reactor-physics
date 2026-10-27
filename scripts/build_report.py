#!/usr/bin/env python3
"""Build a Word (.docx) summary report from the executed notebooks.

Usage:
    python3 scripts/build_report.py [RESULTS_DIR] [OUTPUT.docx]

RESULTS_DIR must contain the four executed notebooks written by run_notebooks.py and
xs_U8_Q1.csv (default: results/ at the repository root). OUTPUT defaults to
RESULTS_DIR/report.docx.

Every number in the report is read by regex from the notebooks' printed output.
If a pattern does not match, the script stops and names the missing number
instead of writing a wrong value. Figures are taken from the image/png outputs
inside the notebooks, so the figures/ folder is not needed.
"""
import base64
import csv
import io
import json
import os
import re
import struct
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DEFAULT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "results")
QS = (1, 2, 3, 4)
NOTEBOOKS = {1: "Q1_pure_U238.ipynb", 2: "Q2_natural_uranium.ipynb",
             3: "Q3_enriched_uranium_light_water.ipynb", 4: "Q4_natural_uranium_heavy_water.ipynb"}
SYSTEM = {1: "Pure U-238",
          2: "Natural uranium",
          3: "2% enriched U + light water",
          4: "Natural U + heavy water"}
EXPECT = {1: "cannot", 2: "cannot", 3: "can", 4: "can"}
MOD = {3: ("H", "H/U", "hydrogen", "H-1"), 4: ("D", "D/U", "deuterium", "D-2")}
# notebook section number -> plot number (one image each)
PLOT_SECTION = {1: 5, 2: 8, 3: 12, 4: 13, 5: 14, 6: 15, 7: 16}
NUM = r"[-+]?\d[\d,]*\.?\d*(?:[eE][-+]?\d+)?"


class ReportError(Exception):
    pass


# ----------------------------------------------------------------------------- parsing
def load_notebook(path):
    """Return {section number: {"text": stdout, "images": [png bytes], "cells": [(source, text)]}}."""
    if not os.path.exists(path):
        raise ReportError(f"notebook not found: {path}")
    nb = json.load(open(path, encoding="utf-8"))
    secs, cur = {}, 0
    for cell in nb["cells"]:
        src = "".join(cell["source"])
        if cell["cell_type"] == "markdown":
            m = re.match(r"\s*###\s+(\d+)\.", src)
            if m:
                cur = int(m.group(1))
            continue
        if cell["cell_type"] != "code":
            continue
        s = secs.setdefault(cur, {"text": "", "images": [], "cells": []})
        text = ""
        for o in cell.get("outputs", []):
            if o.get("output_type") == "stream" and o.get("name") == "stdout":
                t = o["text"]
                text += t if isinstance(t, str) else "".join(t)
            elif o.get("output_type") == "error":
                raise ReportError(f"{path}: a cell in section {cur} raised {o.get('ename')}: {o.get('evalue')}")
            data = o.get("data", {})
            if "image/png" in data:
                png = data["image/png"]
                s["images"].append(base64.b64decode(png if isinstance(png, str) else "".join(png)))
        s["text"] += text
        s["cells"].append((src, text))
    return secs


class Extractor:
    """Regex lookups on one notebook's output; records every value it returns."""

    def __init__(self, q, secs, log):
        self.q, self.secs, self.log = q, secs, log

    def text(self, sec):
        if sec not in self.secs or not self.secs[sec]["text"].strip():
            raise ReportError(f"Q{self.q}: section {sec} has no printed output (notebook not executed?)")
        return self.secs[sec]["text"]

    def get(self, name, sec, pattern, n=1):
        m = re.search(pattern, self.text(sec), re.M)
        if not m:
            raise ReportError(f"Q{self.q}: could not find '{name}' in section {sec} output "
                              f"(pattern {pattern!r})")
        vals = m.groups()[:n] if n > 1 else m.group(1)
        self.log.append((f"Q{self.q}.{name}", vals if n == 1 else " | ".join(vals)))
        return vals

    def all(self, name, sec, pattern, min_count=1):
        rows = re.findall(pattern, self.text(sec), re.M)
        if len(rows) < min_count:
            raise ReportError(f"Q{self.q}: expected at least {min_count} '{name}' lines in section {sec} "
                              f"output, found {len(rows)} (pattern {pattern!r})")
        self.log.append((f"Q{self.q}.{name}", "; ".join(" ".join(r) if isinstance(r, tuple) else r
                                                         for r in rows)))
        return rows


def f(s):
    return float(s.replace(",", ""))


def parse_question(q, secs, log):
    x = Extractor(q, secs, log)
    R = {"q": q}
    R["nu"] = x.get("nu", 1, r"nu=(" + NUM + ")")
    R["E_th_eV"] = x.get("thermal cutoff eV", 1, r"thermal cutoff (" + NUM + r") eV")

    # §3 peak cross sections; §4 thermal spot checks
    R["peak"] = {(nuc, k): v for nuc, k, v in
                 x.all("peak xs (nuclide, reaction, b)", 3,
                       r"^(\S+)\s+MT=\s*\d+ ([eicf]):.*peak (" + NUM + r") b")}
    R["thermal"] = {}
    for nuc, rest in x.all("thermal spot-check lines", 4, r"^\s+(U5|U8|H|D|O)\s+(sigma_.*)$"):
        for k, v in re.findall(r"sigma_([eicf])=\s*(" + NUM + ") b", rest):
            R["thermal"][nuc, k] = v

    # §6 composition
    if q in MOD:
        sym, lab = MOD[q][0], MOD[q][1]
        R["ratio"] = x.get(f"{lab} ratio used", 6, re.escape(lab) + r" atom ratio = (" + NUM + ")")
    comp = x.all("composition", 6, r"^\s+(U5|U8|H|D|O)\s+(" + NUM + r")$")
    R["comp"] = dict(comp)

    # §8 birth spectrum check
    R["birth_mean"], R["birth_exact"], R["p_above_1MeV"] = x.get(
        "birth mean / exact / P(E>1MeV)", 8,
        r"sampled mean (" + NUM + r") MeV vs exact (" + NUM + r")\s+P\(E>1 MeV\) = (" + NUM + ")", 3)

    # §9 slowing-down power per nuclide
    R["xi"] = {nuc: (xi, n) for nuc, xi, n in
               x.all("xi and collisions to thermalise", 9,
                     r"^(\S+)\s+A=.*xi=(" + NUM + r").*\((\d+) collisions to thermalise\)")}

    # §10 analog loop
    R["an_N"], R["an_k"] = x.get("analog N / k", 10, r"analog loop: N=([\d,]+)\s+k_inf=(" + NUM + ")", 2)

    # §11 vectorised run (+ optional 1e9 cell)
    R["N11"], R["time11"], R["rate"], R["fis11"] = x.get(
        "main run N / seconds / neutrons-per-s / fissions", 11,
        r"N = ([\d,]+) in (" + NUM + r")s \(([\d,]+) neutrons/s\)\s+fissions = ([\d,]+)", 4)
    R["k11"], R["e11"] = x.get("main run k / 2sigma", 11,
                               r"k_inf = (" + NUM + r") \+/- (" + NUM + r") \(2 sigma\)", 2)
    R["coll"] = x.get("collisions per history", 11, r"collisions/history = (" + NUM + ")")
    R["hours_1e9"] = x.get("estimated hours for 1e9", 11, r"that is (" + NUM + r") hours at this rate")
    R["err_1e9"] = x.get("estimated 2sigma at 1e9", 11, r"error bar to about \+/-(" + NUM + ")")
    R["big"] = None
    big_cells = [t for src, t in secs.get(11, {"cells": []})["cells"] if "RUN_1E9" in src and "if RUN_1E9" in src]
    if not big_cells:
        raise ReportError(f"Q{q}: no RUN_1E9 cell found in section 11")
    bt = big_cells[0]
    if bt.strip():
        mN = re.search(r"N = ([\d,]+) in (" + NUM + r") h(?: on (\d+) processes)?", bt)
        mk = re.search(r"k_inf = (" + NUM + r") \+/- (" + NUM + r") \(2 sigma\)", bt)
        if not (mN and mk):
            raise ReportError(f"Q{q}: RUN_1E9 cell has output but the N/hours or k_inf line is missing")
        R["big"] = dict(N=mN.group(1), hours=mN.group(2), k=mk.group(1), e=mk.group(2),
                        procs=mN.group(3) or "1")
        log.append((f"Q{q}.1e9 run N / hours / k / 2sigma", " | ".join(R["big"].values())))
    else:
        log.append((f"Q{q}.1e9 run", "not run (RUN_1E9 cell has no output)"))

    # §12 convergence
    R["conv"] = x.all("convergence rows (N, k, 2sigma, rel%)", 12,
                      r"^N=\s*([\d,]+)\s+k_inf=(" + NUM + r") \+/- (" + NUM + r")\s+\(\s*(" + NUM + r")%\)",
                      min_count=2)
    R["k"], R["e"], R["nsig"], R["side"] = x.get(
        "headline k / 2sigma / n-sigma / side", 12,
        r"^k_inf = (" + NUM + r") \+/- (" + NUM + r");\s+(\d+) standard errors (above|below) critical", 4)
    R["N"] = R["conv"][-1][0]

    # §13
    if q in (1, 2):
        R["eta"] = dict(x.all("eta(E) lines", 13, r"^E=\s*(" + NUM + r") MeV\s+eta =\s*(" + NUM + ")",
                              min_count=6))
        R["first_coll"] = x.get("first-collision estimate", 13, r"first-collision estimate .* = (" + NUM + ")")
    else:
        lab = MOD[q][1]
        R["sweep"] = x.all(f"{lab} sweep rows (ratio, k, 2sigma)", 13,
                           re.escape(lab) + r"=\s*(" + NUM + r")\s+k_inf=(" + NUM + r") \+/- (" + NUM + ")",
                           min_count=3)
        R["ok_lo"], R["ok_hi"] = x.get(f"{lab} range with k>1", 13,
                                       r"k_inf > 1 at every ratio tested from " + re.escape(lab)
                                       + r" = (" + NUM + r") to (" + NUM + ")", 2)
        R["edge"] = "edge of the sweep" in x.text(13)
        R["med_fis"] = x.get("median E at fission (MeV)", 13, r"median E at fission = (" + NUM + r") MeV")
        best = max(R["sweep"], key=lambda r: f(r[1]))
        R["best_ratio"], R["best_k"] = best[0], best[1]
        log.append((f"Q{q}.best ratio / k (argmax of sweep rows)", f"{best[0]} | {best[1]}"))

    # §14 slowing down
    R["slow"] = x.all("slowing-down rows (collision, median E, % alive)", 14,
                      r"after\s+(\d+) collisions: median E = (" + NUM + r") MeV, (" + NUM + r")% still alive")

    # §15 lifetimes
    R["t"] = {}
    for g, key in (("fast", "fast"), ("slowing down", "slow"), ("thermal", "th")):
        R["t"][key] = x.get(f"mean time {g} (s, %)", 15,
                            r"mean time " + g + r"\s*: (" + NUM + r") s\s+\(\s*(" + NUM + r")% of the lifetime\)", 2)
    R["life"] = x.get("mean lifetime (s)", 15, r"mean neutron lifetime\s*: (" + NUM + r") s")
    R["reach_th"] = x.get("% reaching thermal", 15, r"neutrons that reach thermal \(< 1 eV\): (" + NUM + r")%")

    # §16 absorption
    R["abs"] = {nuc: (c, fi) for nuc, c, fi in
                x.all("absorption per nuclide (capture, fission)", 16,
                      r"^(\S+)\s+capture (" + NUM + r")\s+fission (" + NUM + r")\s+per source neutron")}
    R["af"], R["am"] = x.get("absorbed in fuel / moderator", 16,
                             r"absorbed in fuel (" + NUM + r")\s+in moderator (" + NUM + ")", 2)
    R["fc_all"] = x.get("fission/capture (all captures)", 16, r"fission-to-capture ratio \(all captures\) = (" + NUM + ")")
    R["fc_fuel"] = x.get("fission/capture (fuel only)", 16, r"fission-to-capture ratio \(fuel only\)\s*= (" + NUM + ")")

    # §17 grid sensitivity
    R["kf17"], R["ef17"] = x.get("full pointwise k / 2sigma", 17,
                                 r"full pointwise data : k_inf = (" + NUM + r") \+/- (" + NUM + ")", 2)
    R["kc17"], R["ec17"], R["pct17"] = x.get("50-point grid k / 2sigma / %", 17,
                                             r"50-point grid only  : k_inf = (" + NUM + r") \+/- (" + NUM
                                             + r")\s+\((" + NUM + r")%\)", 3)

    # §18 conclusion: fission energies
    if q == 1:
        R["med_fis"] = x.get("median E at fission (MeV)", 18, r"Median energy at fission is (" + NUM + r") MeV")
    elif q == 2:
        R["pct_th_fis"] = x.get("% fissions below 1 eV", 18, r"only (" + NUM + r")% of fissions happened below 1 eV")
    else:
        R["pct_th_fis"] = x.get("% fissions below 1 eV", 18, r"(" + NUM + r")% of fissions (?:are now |are )?below 1 eV")

    # figures
    R["img"] = {}
    for p, sec in PLOT_SECTION.items():
        imgs = secs.get(sec, {"images": []})["images"]
        if len(imgs) != 1:
            raise ReportError(f"Q{q}: expected exactly 1 figure (Plot {p}) in section {sec}, found {len(imgs)}")
        R["img"][p] = imgs[0]
    return R


def verdict(k, e):
    k, e = f(k), f(e)
    if k - e > 1:
        return "can go critical"
    if k + e < 1:
        return "cannot go critical"
    return "inconclusive (within 2σ of 1)"


def read_xs_csv(path):
    if not os.path.exists(path):
        raise ReportError(f"U-238 tabulation not found: {path} (it is written by section 4 of {NOTEBOOKS[1]})")
    rows = list(csv.reader(open(path, encoding="utf-8")))
    want = ["E_MeV", "sigma_e", "sigma_i", "sigma_c", "sigma_f"]
    if [h.strip() for h in rows[0]] != want:
        raise ReportError(f"{path}: header is {rows[0]}, expected {want}")
    data = [[float(v) for v in r] for r in rows[1:] if r]
    if len(data) != 50:
        raise ReportError(f"{path}: expected 50 energies, found {len(data)}")
    return data


# ----------------------------------------------------------------------------- docx helpers
def png_size(png):
    if png[:8] != b"\x89PNG\r\n\x1a\n":
        raise ReportError("figure output is not a PNG")
    return struct.unpack(">II", png[16:24])


def rich(par, text, size=None):
    """Add text with tiny markup: _{sub}, ^{sup}, **bold**."""
    for tok in re.split(r"(_\{[^}]*\}|\^\{[^}]*\}|\*\*[^*]+\*\*)", text):
        if not tok:
            continue
        if tok.startswith("_{"):
            run = par.add_run(tok[2:-1]); run.font.subscript = True
        elif tok.startswith("^{"):
            run = par.add_run(tok[2:-1]); run.font.superscript = True
        elif tok.startswith("**"):
            run = par.add_run(tok[2:-2]); run.bold = True
        else:
            run = par.add_run(tok)
        if size:
            run.font.size = Pt(size)
    return par


class Report:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(11)
        for name in ("Heading 1", "Heading 2", "Heading 3"):
            st[name].font.color.rgb = RGBColor(0, 0, 0)
            st[name].font.name = "Calibri"
        for s in self.doc.sections:
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
        self.nfig = 0

    def heading(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def para(self, text):
        return rich(self.doc.add_paragraph(), text)

    def bullets(self, items):
        for t in items:
            rich(self.doc.add_paragraph(style="List Bullet"), t)

    def figure(self, png, caption, note=None):
        w, h = png_size(png)
        width = min(6.3, 4.2 * w / h)            # cap very tall figures at ~4.2 in high
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(io.BytesIO(png), width=Inches(width))
        self.nfig += 1
        c = rich(self.doc.add_paragraph(), f"Figure {self.nfig}. {caption}", size=10)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = bool(note)   # keep figure, caption and explanation together
        if note:
            self.para(note)

    def table(self, header, rows, size=10, widths=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, h in zip(t.rows[0].cells, header):
            cell.text = ""
            rich(cell.paragraphs[0], h, size=size).runs[0].bold = True
        for r in rows:
            cells = t.add_row().cells
            for cell, v in zip(cells, r):
                cell.text = ""
                rich(cell.paragraphs[0], str(v), size=size)
        for row in t.rows:
            for cell in row.cells:
                pf = cell.paragraphs[0].paragraph_format
                pf.space_before = pf.space_after = Pt(0)
        if widths:
            for row in t.rows:
                for cell, wd in zip(row.cells, widths):
                    cell.width = Inches(wd)
        # repeat header row on each page
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
        return t


# ----------------------------------------------------------------------------- content
def sci(v):
    """'2.725e-02' -> '2.725 × 10^{-2}' (rich-text markup); plain numbers pass through."""
    m = re.fullmatch(r"([-+]?\d+\.?\d*)[eE]([-+]?\d+)", v.strip())
    if not m:
        return v
    e = int(m.group(2))
    return m.group(1) if e == 0 else f"{m.group(1)} × 10^{{{e}}}"


def pct_fmt(v):
    return f"{f(v):+.1f}%"


def question_section(rep, R, Q):
    q = R["q"]
    titles = {1: "Q1 — Pure U-238", 2: "Q2 — Natural uranium",
              3: "Q3 — 2% enriched uranium + light water", 4: "Q4 — Natural uranium + heavy water"}
    rep.heading(f"{q + 3}. {titles[q]}", 1)
    side = R["side"]
    v = verdict(R["k"], R["e"])
    if not v.startswith(EXPECT[q]):
        print(f"WARNING: Q{q} verdict '{v}' differs from the expected '{EXPECT[q]} go critical'", file=sys.stderr)
    t = R["t"]

    rep.heading("Results", 2)
    b = [f"k_{{∞}} = **{R['k']} ± {R['e']}** (2σ, N = {R['N']} neutrons) → **{v}**; "
         f"{R['nsig']} standard errors {side} 1.",
         f"Cross-checks: separate vectorised run of {R['N11']} neutrons (used for the tallies below) gives "
         f"{R['k11']} ± {R['e11']}; the one-neutron-at-a-time analog loop ({R['an_N']} neutrons) gives {R['an_k']}.",
         f"Collisions per history: {R['coll']}."]
    if q == 1:
        b.append(f"Fission energies: median {R['med_fis']} MeV — every fission is fast.")
    elif q == 2:
        b.append(f"Fissions below 1 eV: {R['pct_th_fis']}%, although η at thermal is {R['eta']['2.5e-08']} (> 1).")
    else:
        b.append(f"Fissions below 1 eV: {R['pct_th_fis']}% (median fission energy {sci(R['med_fis'])} MeV).")
    b.append(f"Mean neutron lifetime {sci(R['life'])} s: fast {t['fast'][1]}%, slowing down {t['slow'][1]}%, "
             f"thermal {t['th'][1]}%; {R['reach_th']}% of neutrons reach thermal (< 1 eV).")
    b.append(f"Absorptions per source neutron: fuel {R['af']}, moderator {R['am']}.")
    if R["fc_all"] == R["fc_fuel"]:
        b.append(f"Fission-to-capture ratio: {R['fc_all']}.")
    else:
        b.append(f"Fission-to-capture ratio: {R['fc_all']} (all captures), {R['fc_fuel']} (fuel captures only).")
    if q in MOD:
        lab = MOD[q][1]
        b.append(f"Moderator ratio (main runs use {lab} = {R['ratio']}): k_{{∞}} > 1 for every {lab} tested from "
                 f"{R['ok_lo']} to {R['ok_hi']}; best {R['best_k']} at {lab} = {R['best_ratio']}"
                 + (" (upper end is the edge of the sweep)." if R["edge"] else "."))
    b.append(f"Using only the 50-point grid would give {R['kc17']} instead of {R['kf17']} ({pct_fmt(R['pct17'])}).")
    rep.bullets(b)

    rep.heading("Plots", 2)
    # Plot 3: convergence
    conv = R["conv"]
    same_side = all((f(k) - f(e) > 1) if side == "above" else (f(k) + f(e) < 1) for _, k, e, _ in conv)
    rep.figure(R["img"][3], f"Q{q}: k_{{∞}} against number of neutrons tracked (left) and relative error against "
                            f"1/√N (right).",
               f"k_{{∞}} settles at {R['k']} and the relative error falls as 1/√N, from {conv[0][3]}% at "
               f"N = {conv[0][0]} to {conv[-1][3]}% at N = {conv[-1][0]}. "
               + (f"The 2σ band is {side} k_{{∞}} = 1 at every N, so the verdict does not depend on statistics."
                  if same_side else f"At the largest N the 2σ band is clearly {side} k_{{∞}} = 1."))

    # Plot 4
    if q == 1:
        eta = R["eta"]
        rep.figure(R["img"][4], "Q1: energy at fission (left), collisions per history (centre) and "
                                "η(E) = ν σ_{f}/(σ_{f}+σ_{c}) (right).",
                   f"All fissions are fast (median {sci(R['med_fis'])} MeV) because U-238 fission has a threshold near "
                   f"1 MeV. η is {eta['2.0e+00']} at 2 MeV but {eta['1.0e+00']} at 1 MeV and {eta['1.0e-02']} "
                   f"below, so a neutron that slows past the threshold can only be captured, after ~{R['coll']} "
                   f"collisions.")
    elif q == 2:
        eta = R["eta"]
        res = sorted((f(eta["1.0e-04"]), f(eta["1.0e-02"])))
        xi, ncol = R["xi"]["U8"]
        rep.figure(R["img"][4], "Q2: energy at fission (left), collisions per history (centre) and "
                                "η(E) = ν Σ_{f}/(Σ_{f}+Σ_{c}) (right).",
                   f"η at thermal is {eta['2.5e-08']} (above 1), yet only {R['pct_th_fis']}% of fissions occur "
                   f"below 1 eV. Uranium (ξ = {xi}) needs ~{ncol} collisions to thermalise, but capture ends a "
                   f"history after ~{R['coll']}; in the resonance region η is only {res[0]:.3f}–{res[1]:.3f}.")
    else:
        sym, lab, name, _ = MOD[q]
        last = R["sweep"][-1]
        tail = ""
        if R["edge"]:
            tail = f"; the upper limit {R['ok_hi']} is the edge of the sweep, not a bound."
        elif f(last[1]) < f(R["best_k"]) and last[0] != R["best_ratio"]:
            tail = f" and falls to {last[1]} by {lab} = {last[0]} as {name} capture and fuel dilution take over."
        else:
            tail = "."
        rep.figure(R["img"][4], f"Q{q}: energy at fission (left), collisions per history (centre) and k_{{∞}} "
                                f"against the {lab} atom ratio (right).",
                   f"{R['pct_th_fis']}% of fissions are below 1 eV, so the moderator does its job. "
                   f"k_{{∞}} peaks at {R['best_k']} at {lab} = {R['best_ratio']}, is above 1 from {lab} = {R['ok_lo']} "
                   f"to {R['ok_hi']}{tail} This optimum is the expected under-/over-moderation trade-off.")

    # Plot 5: slowing down
    slow = R["slow"]
    first = slow[0]
    th = next((r for r in slow if f(r[1]) < 1e-6), None)
    if q in (1, 2):
        xi = R["xi"]["U8"][0]
        last = slow[-1]
        txt = (f"Each collision with uranium removes very little energy (ξ = {xi}), so the median energy only falls "
               f"from {sci(first[1])} MeV to {sci(last[1])} MeV after {last[0]} collisions, when {last[2]}% of neutrons "
               f"are still alive. The flux stays in the fast and upper slowing-down range and falls off steeply through "
               f"the U-238 resonance region, with no thermal peak, as expected with no moderator.")
    else:
        sym, lab, name, _ = MOD[q]
        xi = R["xi"][sym][0]
        reach = (f"the median energy is below 1 eV after {th[0]} collisions ({sci(th[1])} MeV)" if th else
                 f"the median energy is {sci(slow[-1][1])} MeV after {slow[-1][0]} collisions")
        txt = (f"Collisions with {name} remove a large share of the energy (ξ = {xi}): {reach}. "
               f"The flux shows the fission peak, a roughly flat 1/E slowing-down region with dips at the U-238 "
               f"resonances, and a thermal Maxwellian peak, as expected for a moderated system.")
    rep.figure(R["img"][5], f"Q{q}: neutron energy against collision number (left) and flux per unit lethargy "
                            f"Eφ(E) (right).", txt)

    # Plot 6: lifetimes
    if f(R["reach_th"]) == 0:
        txt = (f"The mean lifetime is {sci(R['life'])} s, spent {t['fast'][1]}% fast and {t['slow'][1]}% slowing down. "
               f"No neutron reaches thermal energy because capture ends every history first.")
    else:
        txt = (f"{R['reach_th']}% of neutrons reach thermal energy, and the thermal group dominates the "
               f"{sci(R['life'])} s mean lifetime ({t['th'][1]}%) because thermal neutrons are slow and travel until "
               f"absorbed.")
        if q == 4 and f(R["life"]) > f(Q[3]["life"]):
            txt += (f" The lifetime is longer than in light water ({sci(Q[3]['life'])} s) since deuterium "
                    f"barely absorbs.")
    rep.figure(R["img"][6], f"Q{q}: mean time per neutron in the fast (> 0.1 MeV), slowing-down "
                            f"(1 eV–0.1 MeV) and thermal (< 1 eV) groups (left) and lifetime distributions (right).",
               txt)

    # Plot 7: absorption
    A = R["abs"]
    if q == 1:
        c8, f8 = A["U-238"]
        txt = (f"With no moderator all absorption is in the fuel: per source neutron U-238 captures {c8} and "
               f"fissions {f8}, a fission-to-capture ratio of {R['fc_all']}.")
    elif q == 2:
        txt = (f"All absorption is in the fuel. U-238 capture ({A['U-238'][0]}) dominates; U-235 adds only "
               f"{A['U-235'][1]} fissions per source neutron, so fission-to-capture is just {R['fc_all']}.")
    else:
        sym, lab, name, nucname = MOD[q]
        txt = (f"The fuel absorbs {R['af']} and the moderator {R['am']} per source neutron ({nucname} capture "
               f"{A[nucname][0]}). U-235 fission ({A['U-235'][1]}) now dominates, raising fission-to-capture "
               f"to {R['fc_all']}.")
    txt += " The Monte Carlo points follow the ratio predicted by the cross sections, confirming the reaction sampling."
    rep.figure(R["img"][7], f"Q{q}: absorptions per source neutron by nuclide (left) and fuel fission-to-capture "
                            f"ratio against energy (right).", txt)


def build(nbdir, out):
    log = []
    Q = {}
    for q in QS:
        secs = load_notebook(os.path.join(nbdir, NOTEBOOKS[q]))
        Q[q] = parse_question(q, secs, log)
    xs = read_xs_csv(os.path.join(nbdir, "xs_U8_Q1.csv"))
    log.append(("xs_U8_Q1.csv rows", str(len(xs))))

    rep = Report()
    d = rep.doc
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Monte Carlo Simulations of Reactor Physics"); r.bold = True; r.font.size = Pt(14)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("k∞ of four infinite homogeneous reactors"); r.bold = True; r.font.size = Pt(13)

    q1, q3, q4 = Q[1], Q[3], Q[4]
    nu = q1["nu"]

    # 1 Problem
    rep.heading("1. Problem", 1)
    rep.bullets([
        "Show by Monte Carlo simulation of neutrons in an infinite homogeneous core that (1) pure U-238 and "
        "(2) natural uranium cannot go critical, while (3) 2% enriched uranium with light water and (4) natural "
        "uranium with heavy water can.",
        "Criterion: the infinite multiplication factor k_{∞} (fission neutrons produced per source neutron). "
        "The system can go critical only if k_{∞} > 1.",
        "Neutrons are born with s(E) = 0.771 E^{1/2} e^{−0.776E} (E in MeV).",
        "Energy-dependent microscopic cross sections (elastic, inelastic, capture, fission), tabulated for U-238 at "
        "(0.2, 0.4, 0.6, 0.8, 1.0) × 10^{y} MeV, y = −8 … 1 (50 energies) as a minimum resolution.",
        "Track 10^{6} and 10^{9} neutrons for good statistics.",
    ])

    # 2 Method
    rep.heading("2. Method", 1)
    rep.bullets([
        "Infinite medium: no leakage, so no geometry or position tracking is needed. Every sampling decision is a "
        "ratio of macroscopic cross sections, so only atom ratios matter, not absolute densities.",
        f"Compositions (atom ratios): Q1 pure U-238; Q2 {f(Q[2]['comp']['U5'])*100:.4g}% U-235; Q3 2% U-235 "
        f"with H/U = {q3['ratio']} as H_{{2}}O; Q4 natural U with D/U = {q4['ratio']} as D_{{2}}O. "
        "The moderator ratio is also swept.",
        "Birth energy: s(E) is exactly a Gamma(3/2, T) distribution with T = 1/0.776 = 1.289 MeV, so it is "
        "sampled directly with no rejection.",
        "Each collision: pick the nuclide with probability N_{i}σ_{t,i}/Σ_{t}; pick the reaction with "
        "probability σ_{x}/σ_{t}. Capture or fission ends the history; scattering continues it.",
        "Elastic: E′ = E[α + (1 − α)ξ], α = ((A − 1)/(A + 1))^{2}, ξ uniform on (0, 1). Inelastic: subtract a "
        "nuclear level sampled from those below E (ENDF MF3 gives only how often, not how much energy is lost).",
        f"Thermal: below {q1['E_th_eV']} eV the energy is resampled from a 293 K Maxwellian instead of slowing "
        "down forever.",
        f"k_{{∞}} = ν × fissions / N with ν = {nu}; σ_{{k}} = ν√(p(1 − p)/N), p = fissions/N. Errors are quoted as 2σ.",
        "Cross sections: pointwise ENDF/B-VIII.0 from KAERI, log-log interpolated. The 50-point tabulation is "
        "written to a CSV per nuclide; the transport uses the full pointwise data because 5 points per decade "
        "cannot resolve the U-238 resonances (section 9 gives the size of that bias).",
    ])

    # 3 Cross sections & verification
    rep.heading("3. Cross sections and verification", 1)
    th = {**q1["thermal"], **Q[2]["thermal"], **q3["thermal"], **q4["thermal"]}
    for need in (("U8", "c"), ("U5", "f"), ("U5", "c"), ("H", "c"), ("D", "c"), ("O", "e")):
        if need not in th:
            raise ReportError(f"thermal spot check sigma_{need[1]} for {need[0]} not found in any notebook")
    log.append(("thermal spot checks used (b)", ", ".join(f"{n}.{k}={th[n, k]}" for n, k in
                                                           (("U8", "c"), ("U5", "f"), ("U5", "c"), ("H", "c"),
                                                            ("D", "c"), ("O", "e")))))
    if ("U8", "c") not in q1["peak"]:
        raise ReportError("Q1: U-238 capture peak not found in section 3 output")
    rep.figure(q1["img"][1], "U-238 microscopic cross sections (ENDF/B-VIII.0 via KAERI). Lines: full pointwise "
                             "data; dots: the 50 tabulated energies.",
               f"Capture shows the resolved resonances (peak {q1['peak']['U8', 'c']} b, Doppler broadened at "
               f"293 K) and fission has a threshold near 1 MeV, as expected for U-238. The coarse dots miss most "
               f"resonance peaks.")
    rep.bullets([
        f"Thermal (0.0253 eV) values match published data: U-238 σ_{{c}} = {th['U8', 'c']} b (ref 2.68), "
        f"U-235 σ_{{f}} = {th['U5', 'f']} b (ref 585), U-235 σ_{{c}} = {th['U5', 'c']} b (ref 99).",
        f"Moderators: H-1 σ_{{c}} = {th['H', 'c']} b (ref 0.3326), D-2 σ_{{c}} = {th['D', 'c']} b (ref 0.000519), "
        f"O-16 σ_{{e}} = {th['O', 'e']} b (ref 3.85).",
    ])
    rep.figure(q1["img"][2], "Sampled birth energies (500,000 neutrons) against s(E), linear (left) and log "
                             "(right) scales.",
               f"The histogram lies on s(E). Sampled mean {q1['birth_mean']} MeV against the exact 3T/2 = "
               f"{q1['birth_exact']} MeV; {f(q1['p_above_1MeV'])*100:.0f}% of neutrons are born above 1 MeV.")
    d.add_page_break()
    rep.para("**Table 1.** U-238 microscopic cross sections (b) at the 50 required energies "
             "(ENDF/B-VIII.0, 4 significant figures).")
    rep.table(["E (MeV)", "σ elastic", "σ inelastic", "σ capture", "σ fission"],
              [[f"{v:.4g}" for v in row] for row in xs], size=8)

    # 4-7 questions
    for q in QS:
        question_section(rep, Q[q], Q)

    # 8 Statistics
    rep.heading("8. Statistics: one million vs one billion neutrons", 1)
    mins = min(int(Q[q]["nsig"]) for q in QS)
    ran = [q for q in QS if Q[q]["big"]]
    notrun = [q for q in QS if not Q[q]["big"]]
    s = []
    for q in ran:
        B = Q[q]["big"]
        s.append(f"Q{q} was run with {B['N']} neutrons ({B['hours']} h on {B['procs']} processes): "
                 f"k_{{∞}} = {B['k']} ± {B['e']} (2σ), "
                 f"against {Q[q]['k']} ± {Q[q]['e']} at N = {Q[q]['N']}.")
    if notrun:
        est = ", ".join(f"Q{q} {Q[q]['hours_1e9']} h" for q in notrun)
        s.append(f"At the measured rates, 10^{{9}} neutrons would take {est}; the error bar would shrink by "
                 f"√1000 ≈ 32.")
    s.append(f"Every verdict is already at least {mins} standard errors from k_{{∞}} = 1, so the extra neutrons add "
             f"precision but change no conclusion.")
    rep.para(" ".join(s))

    # 9 Summary
    rep.heading("9. Summary and limitations", 1)
    rep.table(["Q", "System", "k_{∞} ± 2σ", "Verdict"],
              [[str(q), SYSTEM[q], f"{Q[q]['k']} ± {Q[q]['e']}", verdict(Q[q]["k"], Q[q]["e"])] for q in QS],
              size=10)
    rep.para("")
    pcts = ", ".join(f"Q{q} {pct_fmt(Q[q]['pct17'])}" for q in QS)
    rep.bullets([
        f"Using only the 50-point grid biases k_{{∞}} ({pcts}); it is optimistic where k_{{∞}} > 1, so the full "
        "pointwise data are used for all reported numbers.",
        "Inelastic energy loss is modelled by sampling discrete nuclear levels; the U-238 level list stops at "
        "1.17 MeV, which if anything overestimates Q1's k_{∞}.",
        "Homogeneous mixtures, as the question specifies; a real fuel lattice self-shields the U-238 resonances "
        "and does better.",
        "Q4 assumes pure D_{2}O; any light-water contamination brings back hydrogen capture.",
        f"Scattering is isotropic in the centre of mass, cross sections are at ~293 K, and ν = {nu} is constant.",
    ])

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    d.save(out)
    return log, rep.nfig


def main(argv):
    nbdir = argv[1] if len(argv) > 1 else DEFAULT_DIR
    out = argv[2] if len(argv) > 2 else os.path.join(nbdir, "report.docx")
    try:
        log, nfig = build(nbdir, out)
    except ReportError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1
    print("extracted values:")
    for name, v in log:
        print(f"  {name:55s} {v}")
    print(f"\nwrote {out} ({nfig} figures)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
